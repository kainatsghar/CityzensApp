#!/usr/bin/env python3
"""Generate Cityzens v2.0.0 UAT Sign-Off and Release Notes DOCX files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent

CHANGE_DESCRIPTION = (
    "Deployment of App Version v2.0.0+89 (from production v1.13.2+88). "
    "This is a major release with a full UI/UX redesign across nearly all screens "
    "(login, dashboards, profile, academic modules, fees/challan, attendance, settings, "
    "and supporting flows), plus Academic Hub, PDF Academic Results (view/download/cache), "
    "biometric login with secure credential storage, API-based Launch Promotions "
    "(migrated from Firebase), teacher dashboard caching, parent empty-state when no "
    "linked children, grade-based hiding of Academic Results for Grades 2–4, teacher "
    "edit/delete for diary/classwork/circulars, and related auth, caching, and stability fixes."
)

SPECIAL_INSTRUCTIONS = [
    "Full UI/UX redesign — regression-test all primary Student / Parent / Teacher journeys.",
    "Academic Results hub hidden for Grades 2–4.",
    "Complaint/Query UI remains hidden.",
    "Promotion management screens are debug-only.",
    "See companion full changelog (Cityzens_v2.0.0_Release_Notes) for the complete change list.",
]


def set_cell_shading(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tePr = cell._tc.get_or_add_tcPr()
    cell._tc.get_or_add_tcPr().append(shading)


def style_run(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_bar(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    style_run(run, bold=True, size=14, color=(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def add_label_value_table(doc: Document, rows: list[tuple[str, str]], cols=2) -> None:
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(label)
        style_run(r0, bold=True, size=10)
        r1 = c1.paragraphs[0].add_run(value)
        style_run(r1, size=10)
        set_cell_shading(c0, "D6E3F0")


def checkbox_line(checked: bool, label: str) -> str:
    return f"{'☑' if checked else '☐'}  {label}"


def build_uat_signoff() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("APPENDIX 1-B: UAT Sign-Off Sheet")
    style_run(r, bold=True, size=16, color=(0x1F, 0x4E, 0x79))

    org = doc.add_paragraph()
    org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = org.add_run("Information Technology Department — The City School")
    style_run(r, bold=True, size=11)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("CITYZENS Mobile Application — Version 2.0.0+89")
    style_run(r, size=10, color=(0x66, 0x66, 0x66))

    add_heading_bar(doc, "Header Information")
    add_label_value_table(
        doc,
        [
            ("Change Deployment #", "DEP-Cityzens-v2.0.0"),
            ("Date of First Request", "10-08-2026"),
            ("Requested By", "IT Department (Self Initiated)"),
            ("Department", "IT Department"),
        ],
    )

    add_heading_bar(doc, "SECTION 1: CHANGE REQUEST INFORMATION")
    p = doc.add_paragraph()
    r = p.add_run("Change Type:")
    style_run(r, bold=True, size=11)

    change_types = [
        (True, "Problem Fix"),
        (False, "Query Update"),
        (True, "Parameter addition/update"),
        (False, "Rate Change"),
        (True, "Version Upgrade"),
        (False, "MIS/Report Change"),
        (True, "New Module Addition"),
        (False, "New Product"),
    ]
    for checked, label in change_types:
        para = doc.add_paragraph(checkbox_line(checked, label))
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        for run in para.runs:
            style_run(run, size=10)

    p = doc.add_paragraph()
    r = p.add_run("Change Description:")
    style_run(r, bold=True, size=11)
    desc = doc.add_paragraph(CHANGE_DESCRIPTION)
    for run in desc.runs:
        style_run(run, size=10)
    desc.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    r = p.add_run("Change Area:")
    style_run(r, bold=True, size=11)
    change_areas = [
        (True, "Version Upgrade"),
        (False, "Database"),
        (True, "Application Side"),
        (False, "Parameter Change"),
        (False, "Report"),
        (False, "SMS"),
        (False, "Data Update through backend (Query Attached)"),
        (False, "Other (Specify Here)"),
    ]
    for checked, label in change_areas:
        para = doc.add_paragraph(checkbox_line(checked, label))
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        for run in para.runs:
            style_run(run, size=10)

    add_heading_bar(doc, "Production Roll Out — Implementation Information")
    add_label_value_table(
        doc,
        [
            ("Proposed Deployment Date Time", "10-08-2026 5:30:00 PM"),
            ("Downtime Duration", "Zero Downtime (Rolling Update)"),
            ("Deployment Date/Time and Deployment by", "10-08-2026 5:30:00 PM — IT Software Team"),
            ("By (Name/Role)", "Muhammad Usama (Mobile Application Developer)"),
            ("Application Name", "CITYZENS"),
            ("Server Name/Location", "Production / App Stores (Google Play & Apple App Store)"),
        ],
    )

    add_heading_bar(doc, "Special Instructions / Other Remarks or Cautionary Notes")
    for item in SPECIAL_INSTRUCTIONS:
        para = doc.add_paragraph(item, style="List Bullet")
        for run in para.runs:
            style_run(run, size=10)

    add_heading_bar(doc, "Sign-Off / Approval")
    sig = doc.add_table(rows=5, cols=3)
    sig.style = "Table Grid"
    headers = ["Role", "Name & Signature", "Date"]
    for i, h in enumerate(headers):
        cell = sig.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        style_run(run, bold=True, size=10)
        set_cell_shading(cell, "1F4E79")
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    roles = [
        "UAT By",
        "IT Department – Manager IT Applications",
        "Approved by Head Application Owner",
        "Approved by CTO",
    ]
    for i, role in enumerate(roles, start=1):
        sig.rows[i].cells[0].text = role
        sig.rows[i].cells[1].text = ""
        sig.rows[i].cells[2].text = ""
        for cell in sig.rows[i].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    style_run(run, size=10)

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(18)
    r = footer.add_run(
        "Signatures: Information Technology Department — The City School\n"
        "Companion document: Cityzens_v2.0.0_Release_Notes.docx"
    )
    style_run(r, size=9, color=(0x66, 0x66, 0x66))

    out = OUT_DIR / "Cityzens_App_V2.0.0_UAT_Sign-Off_Sheet.docx"
    doc.save(out)
    return out


RELEASE_NOTES_SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1. App-wide UI/UX redesign (primary callout)",
        [
            "This release is a major visual and interaction overhaul. Nearly all screens have been migrated to the new design system (shared Custom Card, updated theme/colors/typography, refreshed headers/profile cards, shimmer loading states, and tab/chip selectors replacing many dropdowns).",
            "Shared design system: application-wide theme consistency; Custom Card applied as the single card pattern; UI styling tokens updated; color opacity APIs migrated to withValues; Nexa fonts removed.",
            "Auth / onboarding: Login role selection converted from dropdown to tab selector (Student / Teacher / Parent); UserTypeTabSelector modernized; Change Password screen redesigned with particle background and responsive sliver layout; registration / PIN / create-password flows restyled.",
            "Splash / launch: Splash UI refreshed; Launch Promotions presentation updated (API-backed).",
            "Student dashboard & profile: Full student dashboard rebuild; redesigned student info header; enhanced Student Profile Card (roll number, attendance details); updated dashboard cards/widgets with formatted dates and improved layout; financial status UI enhanced.",
            "Parent experience: Separate Parent Dashboard removed; parents use the Student Dashboard with child selector; empty-state UI when no active children are linked; Eye on Your Child component logic/UI refreshed.",
            "Teacher dashboard & profile: Teacher Dashboard rebuilt with scroll-based app bar; Teacher Profile Card with richer info; gender-based avatars (including female avatar); formatted dates and compact layouts; auto section/subject selection; feedback button design refreshed.",
            "Academic Hub (Student & Teacher): New unified hub UI combining Homework, Classwork, and Circulars with shared academic card component.",
            "Academic Results: New PDF results list/viewer/download screens; teacher class-result modern screens restyled; Academic Results hub tile visibility gated by grade.",
            "Attendance / timetable: Attendance loading shimmer; calendar relocated before statistics on student attendance; A-Level subject selection UI improved; teacher attendance/timetable section pages use tab selection instead of dropdowns.",
            "Fees / challan: Updated Challan view and fee challan presentation; fee components reorganized with student info header.",
            "Library & notifications: Library screen and notification/message/challan notification screens restyled under the new theme.",
            "Teacher academic input screens: Add/Edit Diary, Assignment, and Circular screens rebuilt with updated UI components; date pickers constrained to selected date.",
            "Decorative/motion: ParticleBackground performance optimized (compositor-driven); particle density adjusted on selected screens.",
            "Other module screens touched by theme migration include surveys, school events, appointments, OS registration, IEP, and related supporting flows.",
        ],
    ),
    (
        "2. New features",
        [
            "Academic Hub for students and teachers — Homework, Classwork, and Circulars in one tabbed experience.",
            "Academic Results PDFs — list term/archive PDFs, open in-app viewer, download with local caching.",
            "Biometric login — Face ID / fingerprint (device biometrics) with securely stored saved accounts for faster re-login.",
            "My Subjects — student subject list with assigned teachers integrated on the dashboard.",
            "Teacher edit/delete for Homework (Diary), Classwork (Assignment), and Circulars.",
            "Student profile enhancements — roll number, attendance details, richer profile card content.",
            "Gender-based teacher avatars.",
            "App analytics — login and feature engagement tracking on student flows (Firebase Analytics via AppAnalyticsService).",
            "Parent empty-state handling when no linked/active children are available, with refresh path.",
            "Max Marks field updates on class test / marks entry flows.",
            "All PDF Results retrieval/display screen on student dashboard.",
        ],
    ),
    (
        "3. Bug fixes",
        [
            "Promotion / launch promo link handling corrected so taps open the intended URL reliably.",
            "Fee/profile fallback when student profile data is incomplete (avoids broken fee UI).",
            "Teacher date pickers restricted to the currently selected date on diary/assignment/circular entry.",
            "Subject selection disabled during edit of teacher records (prevents wrong section/subject updates); edit resolution improved via section IDs.",
            "Comprehensive logout/provider state cleanup so prior user data does not linger across sessions/accounts.",
            "FeeChallanProvider synced on child switch.",
            "“Late” removed from student dashboard attendance service object usage (attendance display cleanup).",
            "API authorization header handling improved for authenticated requests.",
        ],
    ),
    (
        "4. Performance / caching",
        [
            "Teacher cache: profile, sections, and subjects loaded once at dashboard init and reused across screens (refreshed on app restart).",
            "Per-student profile caching for faster child/dashboard loads.",
            "Results PDF disk cache to avoid re-downloading the same PDFs.",
            "Student dashboard force-refresh / class filtering for results data freshness.",
            "Dashboard initialization moved toward repository pattern; standardized clearFormState cleanup.",
            "Particle background performance optimization.",
        ],
    ),
    (
        "5. Security / auth",
        [
            "Biometric authentication service (local_auth).",
            "Secure credential storage (flutter_secure_storage) for saved accounts.",
            "Session/provider wipe on logout across major modules.",
            "Improved API authorization header logic.",
            "Password updates can sync into saved biometric accounts where applicable.",
        ],
    ),
    (
        "6. API / backend integration",
        [
            "Launch Promotions migrated from Firebase/Firestore to ERP API (GetAllLaunchPromotions and related create/update/delete); launch promo still shown between splash and login.",
            "Results APIs: PDFResults / AllPDFResults plus PDF download/open pipeline.",
            "Student Subjects API wired into dashboard.",
            "Teacher CRUD APIs for update/delete of diary, assignment, and circular/announcement records.",
            "Centralized URL launching logic; Android manifest queries updated for improved link handling.",
            "Circular file attachment URL source updated.",
        ],
    ),
    (
        "7. Hidden / disabled / grade-gated behavior",
        [
            "Academic Results hub hidden for Grades 2, 3, and 4 (Grades 5–6 currently re-enabled for Results).",
            "Complaint/Query UI remains hidden/disabled in the student hub (code retained; not user-facing for this release).",
            "Dashboard hub tiles still present but inactive (onTap null): Events, Survey, Appointment, Withdrawal, Complaints.",
            "Promotion management screens available in debug builds only (not for production end users).",
        ],
    ),
    (
        "8. Platform / dependency upgrades",
        [
            "App version bumped from 1.13.2+88 to 2.0.0+89.",
            "Gradle upgraded to 8.14; iOS CocoaPods re-linked.",
            "Added: local_auth, flutter_secure_storage, connectivity_plus, path_provider, open_filex, flutter_inappwebview.",
            "Firebase packages upgraded (core, messaging, crashlytics, remote_config, analytics, cloud_firestore).",
            "Other dependency bumps: image_picker/cropper, device_info_plus, package_info_plus, flutter_local_notifications, file_picker, upgrader, animate_do, lottie, etc.",
            "open_file replaced with open_filex; grouped_list / introduction_screen removed from pubspec in this line.",
        ],
    ),
    (
        "9. Suggested UAT focus",
        [
            "Login: role tabs + biometric save / login / remove account for Student, Teacher, and Parent.",
            "Full visual regression of Student Dashboard, Teacher Dashboard, profiles, Academic Hub, attendance, challan, and change password.",
            "Academic Hub tabs + teacher edit/delete for diary, assignment, circular.",
            "Academic Results PDF list / open / download; confirm Grades 2–4 do not see Results hub.",
            "Parent with no linked children → empty state + refresh.",
            "Launch promotion display and external link open.",
            "Teacher mark attendance / timetable using cached sections & subjects; auto-selection behavior.",
            "Fee challan after child switch.",
            "Logout then login as another user (no data bleed).",
        ],
    ),
    (
        "10. Branch / version reference",
        [
            "Production (live): branch version-1.13.2+88 — app version 1.13.2+88.",
            "Latest (this release): branch new_temp_v_1_13_3 — app version 2.0.0+89.",
            "Approx. 49 commits between branches; large application-side UI and feature delta.",
            "UAT Sign-Off reference: DEP-Cityzens-v2.0.0 — Cityzens_App_V2.0.0_UAT_Sign-Off_Sheet.docx.",
        ],
    ),
    (
        "11. Internal / non-user-facing (optional awareness)",
        [
            "Technical documentation vault initialized under docs/.",
            "Postman API collection and Live/Test environment configs added under postman/.",
        ],
    ),
]


def build_release_notes_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("CITYZENS — Release Notes / Full Changelog")
    style_run(r, bold=True, size=16, color=(0x1F, 0x4E, 0x79))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Version 2.0.0+89  |  From production 1.13.2+88\n"
        "Branches: new_temp_v_1_13_3 ← version-1.13.2+88\n"
        "Document date: 10-08-2026"
    )
    style_run(r, size=10, color=(0x66, 0x66, 0x66))

    intro = doc.add_paragraph(
        "This document lists the full product-facing changes between the live production "
        "branch and the latest release candidate. Lead change: full UI/UX redesign across "
        "nearly all screens, plus new Academic Hub, PDF Results, biometric login, and "
        "API-based promotions."
    )
    for run in intro.runs:
        style_run(run, size=10)

    for heading, bullets in RELEASE_NOTES_SECTIONS:
        add_heading_bar(doc, heading)
        for bullet in bullets:
            para = doc.add_paragraph(bullet, style="List Bullet")
            for run in para.runs:
                style_run(run, size=10)

    out = OUT_DIR / "Cityzens_v2.0.0_Release_Notes.docx"
    doc.save(out)
    return out


def build_release_notes_md() -> Path:
    lines = [
        "# CITYZENS — Release Notes / Full Changelog",
        "",
        "**Version:** 2.0.0+89 (from production **1.13.2+88**)",
        "",
        "**Branches:** `new_temp_v_1_13_3` ← `version-1.13.2+88`",
        "",
        "**Document date:** 10-08-2026",
        "",
        "This document lists the full product-facing changes between the live production "
        "branch and the latest release candidate. Lead change: **full UI/UX redesign across "
        "nearly all screens**, plus new Academic Hub, PDF Results, biometric login, and "
        "API-based promotions.",
        "",
        "Companion UAT form: [Cityzens_App_V2.0.0_UAT_Sign-Off_Sheet.docx](Cityzens_App_V2.0.0_UAT_Sign-Off_Sheet.docx)",
        "",
    ]
    for heading, bullets in RELEASE_NOTES_SECTIONS:
        lines.append(f"## {heading}")
        lines.append("")
        for bullet in bullets:
            lines.append(f"- {bullet}")
        lines.append("")
    out = OUT_DIR / "Cityzens_v2.0.0_Release_Notes.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


if __name__ == "__main__":
    uat = build_uat_signoff()
    md = build_release_notes_md()
    notes = build_release_notes_docx()
    print(uat)
    print(md)
    print(notes)
